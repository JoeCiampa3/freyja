"""
Freyja pre-processor
====================
Builds models/freyja.xml (the MuJoCo model) from models/freyja_template.xml by
pulling numbers out of the Freyja Anthropometric Reference Google Sheet.

    ground-truth sheet  -->  this script  -->  freyja.xml
    (named ranges)           (check + fill)     (never edit by hand!)

HOW TO RUN (from the freyja-sim folder, using the project's virtual environment)
    freyja.venv\\Scripts\\python scripts\\pre_processor.py
    freyja.venv\\Scripts\\python scripts\\pre_processor.py --xlsx "C:\\path\\to\\sheet.xlsx"
        works offline from a downloaded copy of the sheet (needs: pip install openpyxl)
    freyja.venv\\Scripts\\python scripts\\pre_processor.py --strict    fail if any ROM is missing
    freyja.venv\\Scripts\\python scripts\\pre_processor.py --dry-run   check everything, write nothing

WHAT IT PULLS (three named ranges in the sheet: Data > Named ranges)
    bsip : mass, length, CoM, inertia tensor per segment
    pos  : body positions (pos_x/y/z) and joint anchors (jpos_z) per segment
    rom  : joint range-of-motion limits (the "MuJoCo limit - FINAL" column)
    Segments are looked up by NAME and columns by HEADER, never by row/column
    number, so inserting rows or reordering columns in the sheet is safe.
    Renaming a segment or a header is NOT safe: the script will say so.

HOW TO WRITE PLACEHOLDERS IN THE TEMPLATE
    ${thigh_right_mass}  or  !{thigh_right_mass}      (both styles work)
        BSIP/position names are  <segment>_<value>, e.g.
        pelvis_com_x   thigh_left_ixy   upper_arm_right_pos_z   abdomen_jpos_z
        Bilateral segments (upper_arm forearm hand thigh shank foot) exist as
        _right (straight from the sheet) and _left (mirrored in y automatically).
    !{hip_flexion}   ROM names are  <joint>_<movement>  exactly as in the sheet,
        lower-case with underscores: hip_internal_rotation, ankle_dorsiflexion ...
    !{-hip_extension}   a leading minus negates the value. ROM values are stored
        as positive magnitudes, so the template decides the sign. Because each
        joint axis already fixes which movement is positive, the sign convention
        sits right next to the axis it belongs to:
            axis="0 -1 0" range="!{-hip_extension} !{hip_flexion}"

WHEN A VALUE IS MISSING
    BSIP / position values:  always an error (there is no sane default).
    ROM values:             the joint is left unlimited, a FIXME comment is put in
                            the XML and a warning is printed (same thing you did
                            by hand for the spine). Use --strict to make it fatal.

EVERY RUN ALSO
    * checks the numbers (formats, positive mass, valid inertia tensors, chain
      lengths, ROM sanity) and reports ALL problems at once with the sheet cell
    * checks the finished XML is well-formed and that MuJoCo can compile it
    * never replaces a good freyja.xml with a broken one
    * saves a CSV of every value used to archive/ (only when something changed)
"""

from __future__ import annotations

import argparse
import csv
import difflib
import io
import math
import os
import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np

# =============================================================================
# SETTINGS  (the only section you should normally need to touch)
# =============================================================================
HERE = Path(__file__).resolve().parent
ROOT = HERE.parent  # the freyja-sim folder

SHEET_KEY = "16-XKTGIO4FvCfACWS000RBR2BNFJ0ho6tvMeRswaV2Y"
CREDENTIALS = Path(os.environ.get(  # set FREYJA_CREDENTIALS to use a different key file
    "FREYJA_CREDENTIALS",
    ROOT.parents[1] / "Misc" / "project-valkyrie-509701-0b68d6f5e84d.json"))
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]  # read-only

TEMPLATE = ROOT / "models" / "freyja_template.xml"
OUTPUT = ROOT / "models" / "freyja.xml"
FAILED_OUTPUT = ROOT / "models" / "freyja_FAILED.xml"  # a broken build is parked here
ARCHIVE_DIR = ROOT / "archive"

# Names of the named ranges in the Google Sheet
RANGE_BSIP, RANGE_POS, RANGE_ROM = "bsip", "pos", "rom"

# The sheet must contain exactly these segments (after ident()). If someone adds,
# renames or deletes one in the sheet the script stops and says so ("sheet drift").
EXPECTED_SEGMENTS = ["head_neck", "thorax", "abdomen", "pelvis", "upper_arm",
                     "forearm", "hand", "thigh", "shank", "foot"]
BILATERAL = ["upper_arm", "forearm", "hand", "thigh", "shank", "foot"]

# Column headers (units in brackets are ignored: "mass (kg)" -> "mass")
BSIP_COLUMNS = ["mass", "length", "com_x", "com_y", "com_z",
                "ixx", "iyy", "izz", "ixy", "ixz", "iyz"]
POS_COLUMNS = ["pos_x", "pos_y", "pos_z", "jpos_z"]

# Mirroring the right side into the left (MuJoCo +y = left). Only values that
# involve y change sign: CoM y, the products of inertia containing y, body pos y.
FLIP_BSIP = {"com_y", "ixy", "iyz"}
FLIP_POS = {"pos_y"}

# In the ROM table, which column holds the value MuJoCo should use
ROM_VALUE_HEADER_KEYWORD = "final"  # matches "MuJoCo limit - FINAL (deg)"
ROM_MAX_DEG = 360.0

# Each body sits one parent-segment-length from its parent: (child, segment whose
# length it should equal). A mismatch only warns, since you may do it on purpose.
CHAIN_CHECKS = [("abdomen", "abdomen"), ("thorax", "thorax"),
                ("forearm_right", "upper_arm_right"), ("hand_right", "forearm_right"),
                ("shank_right", "thigh_right"), ("foot_right", "shank_right")]
CHAIN_TOLERANCE_M = 0.001

# =============================================================================
# SMALL HELPERS
# =============================================================================


def ident(name) -> str:
    """'Head & Neck' -> 'head_neck'. Turns any sheet label into a safe key."""
    return re.sub(r"\W+", "_", str(name).strip().lower()).strip("_")


def column_key(header) -> str:
    """'CoM_x (m)' -> 'com_x', 'mass (kg)' -> 'mass'. Bracketed units are dropped."""
    if header is None:
        return ""
    return ident(re.sub(r"\(.*?\)", "", str(header)))


def fmt(v) -> str:
    """Number -> text for the XML. Also removes the -0.0 ambiguity."""
    x = float(v)
    if x == 0.0:
        x = 0.0
    return f"{x:.12g}"


def to_number(value):
    """Cell -> float, or None if the cell is blank. Raises ValueError if it is
    text, an error like #REF!, NaN or infinity."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    x = float(value)  # ValueError for text
    if not math.isfinite(x):
        raise ValueError(value)
    return x


def col_letter(n: int) -> str:
    """1 -> A, 27 -> AA"""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


class Report:
    """Collects every problem so you see them all at once, not one per run."""

    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, msg: str):
        self.errors.append(msg)

    def warn(self, msg: str):
        self.warnings.append(msg)

    def stop_if_errors(self, stage: str):
        if not self.errors:
            return
        print(f"\nStopped: {len(self.errors)} problem(s) found while {stage}")
        for e in self.errors:
            print(f"  ERROR   {e}")
        for w in self.warnings:
            print(f"  WARNING {w}")
        sys.exit(1)


# =============================================================================
# 1. GETTING THE RAW TABLES (Google Sheet, or a downloaded .xlsx for offline use)
# =============================================================================


@dataclass
class RawTable:
    """One named range exactly as it came from the sheet (header row first)."""
    name: str
    sheet: str
    first_row: int  # sheet row number of rows[0]
    first_col: int  # sheet column number of rows[0][0]
    rows: list

    @classmethod
    def from_range(cls, name, range_str, rows):
        m = re.match(r"^'?(.+?)'?!\$?([A-Z]+)\$?(\d+)", range_str)
        if not m:
            raise ValueError(f"cannot understand range '{range_str}'")
        col = 0
        for ch in m.group(2):
            col = col * 26 + ord(ch) - 64
        return cls(name, m.group(1), int(m.group(3)), col, rows)

    def ref(self, row_idx: int, col_idx: int) -> str:
        """Where a cell lives in the sheet, e.g. 'MuJoCo Reference'!G20"""
        sheet = f"'{self.sheet}'" if " " in self.sheet else self.sheet
        return f"{sheet}!{col_letter(self.first_col + col_idx)}{self.first_row + row_idx}"


def fetch_google() -> dict:
    import gspread  # imported here so --xlsx and the tests work without it
    from google.oauth2.service_account import Credentials

    if not CREDENTIALS.exists():
        sys.exit(f"Credentials file not found:\n  {CREDENTIALS}\n"
                 "Fix the path in the SETTINGS section, or set FREYJA_CREDENTIALS.")
    creds = Credentials.from_service_account_file(str(CREDENTIALS), scopes=SCOPES)
    try:
        sh = gspread.authorize(creds).open_by_key(SHEET_KEY)
        resp = sh.values_batch_get([RANGE_BSIP, RANGE_POS, RANGE_ROM],
                                   params={"valueRenderOption": "UNFORMATTED_VALUE"})
    except gspread.exceptions.SpreadsheetNotFound:
        sys.exit("Google says the sheet doesn't exist or isn't shared with this account.\n"
                 f"Share it (Viewer is enough) with: {creds.service_account_email}")
    except gspread.exceptions.APIError as e:
        status = getattr(getattr(e, "response", None), "status_code", "?")
        hint = {400: f"One of the named ranges ({RANGE_BSIP}, {RANGE_POS}, {RANGE_ROM}) "
                     "was not found. Check Data > Named ranges in the sheet.",
                403: f"Not allowed. Share the sheet with: {creds.service_account_email}",
                404: "Sheet not found. Check SHEET_KEY."}.get(status, "")
        sys.exit(f"Google Sheets API error {status}. {hint}\n{e}")
    except Exception as e:  # network down, DNS, timeouts ...
        sys.exit(f"Could not reach Google Sheets ({type(e).__name__}: {e}).\n"
                 "No internet? You can work offline with --xlsx <downloaded sheet>.")

    tables = {}
    for name, vr in zip((RANGE_BSIP, RANGE_POS, RANGE_ROM), resp["valueRanges"]):
        tables[name] = RawTable.from_range(name, vr["range"], vr.get("values", []))
    return tables


def fetch_xlsx(path: str) -> dict:
    try:
        import openpyxl
    except ImportError:
        sys.exit("Reading an .xlsx needs openpyxl:  freyja.venv\\Scripts\\pip install openpyxl")
    if not Path(path).exists():
        sys.exit(f"Spreadsheet file not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)  # data_only = the computed values
    tables = {}
    for name in (RANGE_BSIP, RANGE_POS, RANGE_ROM):
        if name not in wb.defined_names:
            sys.exit(f"Named range '{name}' not found in {path}. "
                     f"Found: {list(wb.defined_names.keys())}")
        dest = wb.defined_names[name].attr_text  # 'MuJoCo Reference'!$A$14:$L$24
        sheet, rng = dest.rsplit("!", 1)
        ws = wb[sheet.strip("'")]
        rows = [["" if c.value is None else c.value for c in row]
                for row in ws[rng.replace("$", "")]]
        tables[name] = RawTable.from_range(name, f"{sheet}!{rng}", rows)
    return tables


# =============================================================================
# 2. TURNING TABLES INTO NUMBERS (and checking them)
# =============================================================================


@dataclass
class Param:
    value: float
    source: str  # where it came from, for error messages and the CSV archive
    table: str = ""  # bsip / pos / rom


def read_segment_table(raw: RawTable, columns: list, report: Report) -> dict:
    """bsip / pos table -> {segment: {column: Param}}. Checks segment names,
    headers and that every needed cell is a real number."""
    if not raw.rows:
        report.error(f"[{raw.name}] the named range is empty")
        return {}
    header = [column_key(h) for h in raw.rows[0]]
    if "segment" not in header:
        report.error(f"[{raw.name}] no 'Segment' header found. Headers seen: {raw.rows[0]}")
        return {}
    seg_i = header.index("segment")

    col_i = {}
    for col in columns:
        if col in header:
            col_i[col] = header.index(col)
        else:
            report.error(f"[{raw.name}] column '{col}' not found. "
                         f"Headers seen: {[h for h in header if h]}")

    found = {}
    for r, row in enumerate(raw.rows[1:], start=1):
        label = row[seg_i] if seg_i < len(row) else ""
        if str(label).strip() == "":
            continue
        seg = ident(label)
        if seg in found:
            report.error(f"[{raw.name}] segment '{label}' appears twice "
                         f"(second one at {raw.ref(r, seg_i)})")
            continue
        params = {}
        for col, ci in col_i.items():
            cell = row[ci] if ci < len(row) else ""
            ref = raw.ref(r, ci)
            try:
                x = to_number(cell)
            except (ValueError, TypeError):
                report.error(f"[{raw.name}] {ref} ({label} / {col}): {cell!r} is not a number")
                continue
            if x is None:
                report.error(f"[{raw.name}] {ref} ({label} / {col}): cell is empty")
                continue
            params[col] = Param(x, ref, raw.name)
        found[seg] = params

    missing = [s for s in EXPECTED_SEGMENTS if s not in found]
    extra = [s for s in found if s not in EXPECTED_SEGMENTS]
    if missing or extra:
        report.error(f"[{raw.name}] sheet drift. Missing segments: {missing or 'none'}; "
                     f"unexpected segments: {extra or 'none'}. Recent sheet edit vs "
                     "EXPECTED_SEGMENTS in pre_processor.py?")
    return found


def read_rom_table(raw: RawTable, report: Report):
    """rom table -> ({'hip_flexion': 133.8, ...}, {names that exist but are blank})"""
    values: dict[str, Param] = {}
    blank: set[str] = set()
    if not raw.rows:
        report.error(f"[{raw.name}] the named range is empty")
        return values, blank
    header = [column_key(h) for h in raw.rows[0]]
    for needed in ("joint", "movement"):
        if needed not in header:
            report.error(f"[{raw.name}] no '{needed}' header found. Headers seen: {raw.rows[0]}")
            return values, blank
    final = [i for i, h in enumerate(header) if ROM_VALUE_HEADER_KEYWORD in h]
    if len(final) != 1:
        report.error(f"[{raw.name}] expected exactly one column whose header contains "
                     f"'{ROM_VALUE_HEADER_KEYWORD}', found {len(final)}. Headers: {raw.rows[0]}")
        return values, blank
    ji, mi, vi = header.index("joint"), header.index("movement"), final[0]

    functional_i = header.index("functional") if "functional" in header else None
    for r, row in enumerate(raw.rows[1:], start=1):
        cells = lambda i: row[i] if i < len(row) else ""  # noqa: E731
        if str(cells(ji)).strip() == "" or str(cells(mi)).strip() == "":
            continue
        key = ident(f"{cells(ji)} {cells(mi)}")
        if key in values or key in blank:
            report.error(f"[{raw.name}] '{cells(ji)} / {cells(mi)}' appears twice "
                         f"(second one at {raw.ref(r, mi)})")
            continue
        ref = raw.ref(r, vi)
        try:
            x = to_number(cells(vi))
        except (ValueError, TypeError):
            report.error(f"[{raw.name}] {ref} ({key}): {cells(vi)!r} is not a number")
            continue
        if x is None:
            blank.add(key)
            continue
        if not 0 <= x <= ROM_MAX_DEG:
            report.error(f"[{raw.name}] {ref} ({key}): {x:g} deg is outside 0..{ROM_MAX_DEG:g}. "
                         "ROM values are positive magnitudes; the template supplies the sign.")
            continue
        values[key] = Param(x, ref, raw.name)
        if functional_i is not None:  # soft check: can the limit even do the task?
            try:
                demand = to_number(cells(functional_i))
            except (ValueError, TypeError):
                demand = None  # the column also holds notes like 'low' or '~0'
            if demand is not None and demand > x:
                report.warn(f"[{raw.name}] {key}: MuJoCo limit {x:g} deg is below the "
                            f"functional demand {demand:g} deg")
    return values, blank


def check_bsip(bsip: dict, report: Report):
    """Physical sanity of each segment's mass properties (before mirroring)."""
    for seg, p in bsip.items():
        if len(p) < len(BSIP_COLUMNS):
            continue  # a missing/bad cell was already reported
        v = {k: x.value for k, x in p.items()}
        where = f"[bsip] {seg} (row source {p['mass'].source})"
        if v["mass"] <= 0:
            report.error(f"{where}: mass must be > 0, got {v['mass']:g}")
        if v["length"] <= 0:
            report.error(f"{where}: length must be > 0, got {v['length']:g}")
        I = np.array([[v["ixx"], v["ixy"], v["ixz"]],
                      [v["ixy"], v["iyy"], v["iyz"]],
                      [v["ixz"], v["iyz"], v["izz"]]])
        a, b, c = np.linalg.eigvalsh(I)  # principal moments, ascending
        if a <= 0:
            report.error(f"{where}: inertia tensor is not positive definite "
                         f"(smallest principal moment {a:.3g}). Check the signs of the products of inertia.")
        elif a + b < c * (1 - 1e-9):
            report.error(f"{where}: principal moments {a:.4g}, {b:.4g}, {c:.4g} break the triangle "
                         "inequality (a + b >= c). No real object has this inertia and MuJoCo will refuse it.")
        com = math.sqrt(v["com_x"] ** 2 + v["com_y"] ** 2 + v["com_z"] ** 2)
        if v["length"] > 0 and com > v["length"]:
            report.warn(f"{where}: centre of mass is {com:.3f} m from the origin, "
                        f"further than the segment is long ({v['length']:.3f} m)")


def mirror(params: dict, flip: set) -> dict:
    """Bilateral segments become <seg>_right (as in the sheet) and <seg>_left
    (mirrored in y). Everything else is passed through unchanged."""
    out = {}
    for seg, p in params.items():
        if seg in BILATERAL:
            out[f"{seg}_right"] = dict(p)
            out[f"{seg}_left"] = {k: Param(-q.value if k in flip else q.value,
                                           q.source + " (mirrored)", q.table)
                                  for k, q in p.items()}
        else:
            out[seg] = p
    return out


def flatten(params: dict) -> dict:
    """{'thigh_right': {'mass': Param}} -> {'thigh_right_mass': Param}"""
    return {f"{seg}_{key}": q for seg, p in params.items() for key, q in p.items()}


def check_positions(values: dict, report: Report) -> float | None:
    """Checks on the (mirrored) numbers. Returns the standing ankle height (m)."""
    v = {k: p.value for k, p in values.items()}
    for k, x in v.items():
        if "pos" in k and abs(x) > 2.0:
            report.error(f"[pos] {k} = {x:g} m looks wrong (units? expected metres)")

    for child, parent in CHAIN_CHECKS:
        try:
            gap = abs(v[f"{child}_pos_z"]) - v[f"{parent}_length"]
        except KeyError:
            continue
        if abs(gap) > CHAIN_TOLERANCE_M:
            report.warn(f"[pos] {child} sits {abs(v[f'{child}_pos_z']):.4f} m from its parent but "
                        f"{parent} is {v[f'{parent}_length']:.4f} m long (off by {gap * 1000:+.1f} mm)")

    for k in ("thigh_right_pos_y", "upper_arm_right_pos_y"):
        if k in v and v[k] >= 0:
            report.warn(f"[pos] {k} = {v[k]:g}: right-side bodies should have NEGATIVE y "
                        "(MuJoCo +y is the subject's left)")

    try:
        ankle = sum(v[f"{s}_pos_z"] for s in ("pelvis", "thigh_right", "shank_right", "foot_right"))
    except KeyError:
        return None
    if ankle < 0:
        report.error(f"[pos] standing chain puts the ankle {ankle * 1000:.1f} mm BELOW the floor "
                     "(pelvis + thigh + shank + foot heights should sum to the ankle height)")
    elif ankle > 0.15:
        report.warn(f"[pos] ankle is {ankle * 1000:.0f} mm above the floor in the standing pose; expected ~60-70 mm")
    return ankle


# =============================================================================
# 3. FILLING IN THE TEMPLATE
# =============================================================================

# ${name} and !{name} are the same thing; a leading '-' inside negates the value.
ANY_BRACES = re.compile(r"([$!])\{([^}]*)\}")
VALID_NAME = re.compile(r"^\s*(-?)\s*([A-Za-z_][A-Za-z0-9_]*)\s*$")
JOINT_TAG = re.compile(r"<joint\b[^<>]*>")
RANGE_ATTR = re.compile(r'\s+range="([^"]*)"')


@dataclass
class Placeholder:
    start: int
    end: int
    line: int
    negate: bool
    key: str


def find_placeholders(text: str, report: Report | None = None) -> list:
    found = []
    for m in ANY_BRACES.finditer(text):
        line = text.count("\n", 0, m.start()) + 1
        ok = VALID_NAME.match(m.group(2))
        if not ok:
            if report is not None:
                report.error(f"template line {line}: placeholder '{m.group(0)}' has no valid name. "
                             "Put a value name inside, e.g. !{hip_flexion}")
            continue
        found.append(Placeholder(m.start(), m.end(), line, ok.group(1) == "-", ok.group(2)))
    return found


def unlimit_joints(text: str, missing: set, report: Report) -> str:
    """For every <joint> whose range= uses a value that is missing: drop
    limited/range and leave a FIXME comment (what you did by hand for the spine)."""

    def fix(m):
        tag = m.group(0)
        rng = RANGE_ATTR.search(tag)
        if not rng:
            return tag
        gone = sorted({p.key for p in find_placeholders(rng.group(1)) if p.key in missing})
        if not gone:
            return tag
        name = re.search(r'name="([^"]*)"', tag)
        name = name.group(1) if name else "?"
        tag = re.sub(r'\s+limited="true"', "", RANGE_ATTR.sub("", tag))
        report.warn(f"ROM missing for joint '{name}' (template line "
                    f"{text.count(chr(10), 0, m.start()) + 1}): {', '.join(gone)}. "
                    "Joint left UNLIMITED until the sheet has the value.")
        return f"{tag} <!-- FIXME(pre_processor): no ROM in sheet for {', '.join(gone)}; joint left unlimited -->"

    return JOINT_TAG.sub(fix, text)


def render(text: str, values: dict, report: Report, strict: bool = False) -> str:
    """Replace every placeholder. Unknown names are collected into the report."""
    phs = find_placeholders(text, report)
    missing = {p.key for p in phs if p.key not in values}
    if missing and not strict:
        text = unlimit_joints(text, missing, report)
        phs = find_placeholders(text, report)
        missing = {p.key for p in phs if p.key not in values}

    for p in phs:
        if p.key in missing:
            hint = difflib.get_close_matches(p.key, values.keys(), n=3, cutoff=0.7)
            report.error(f"template line {p.line}: '{p.key}' is not in the sheet data"
                         + (f". Did you mean: {', '.join(hint)}?" if hint else ""))
    if report.errors:
        return text

    def sub(m):
        ok = VALID_NAME.match(m.group(2))
        x = values[ok.group(2)]
        return fmt(-x if ok.group(1) == "-" else x)

    return ANY_BRACES.sub(sub, text)


def warn_unused(text: str, values: dict, report: Report):
    """Flag sheet data the template ignores, so a new sheet value can't silently
    go nowhere. (Unused BSIP values like 'length' are normal; position and ROM
    values are not, unless a position is zero.)"""
    used = {p.key for p in find_placeholders(text)}
    for k, p in sorted(values.items()):
        if k in used or p.table == "bsip":
            continue
        if p.table == "pos" and abs(p.value) < 1e-9:
            continue
        report.warn(f"sheet value {k} = {fmt(p.value)} ({p.source}) is not used by the template")


def add_banner(xml: str) -> str:
    banner = (f"<!-- GENERATED by scripts/pre_processor.py from models/{TEMPLATE.name}. "
              "Do not edit by hand: change the template or the Google Sheet and re-run. -->")
    decl = re.match(r"\s*<\?xml[^>]*\?>\s*", xml)
    if decl:  # a comment may not come before the XML declaration
        return xml[:decl.end()] + banner + "\n" + xml[decl.end():]
    return banner + "\n" + xml


# =============================================================================
# 4. CHECKING THE RESULT
# =============================================================================


def check_xml(xml: str, report: Report):
    """Well-formed? Do the joint ranges make sense?"""
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as e:
        line = xml.splitlines()[e.position[0] - 1].strip() if e.position[0] <= len(xml.splitlines()) else ""
        report.error(f"generated XML is not well-formed: {e}\n          -> {line}")
        return
    for j in root.iter("joint"):
        rng = j.get("range")
        if rng is None:
            continue
        try:
            lo, hi = (float(t) for t in rng.split())
        except ValueError:
            report.error(f"joint '{j.get('name')}': range='{rng}' is not two numbers")
            continue
        if lo >= hi:
            report.error(f"joint '{j.get('name')}': range min {lo:g} is not below max {hi:g}")
        elif lo > 0 or hi < 0:
            report.warn(f"joint '{j.get('name')}': range {lo:g}..{hi:g} excludes the zero pose")


def compile_with_mujoco(xml: str, report: Report):
    """The real test: can MuJoCo load it? Returns the model or None."""
    try:
        import mujoco
    except ImportError:
        report.warn("mujoco is not installed here, so the compile check was skipped")
        return None
    try:
        return mujoco.MjModel.from_xml_string(xml)
    except Exception as e:  # mujoco raises plain ValueError with its own message
        report.error(f"MuJoCo could not compile the model: {e}")
        return None


# =============================================================================
# 5. ARCHIVE
# =============================================================================


def archive_values(values: dict, directory: Path) -> str:
    """Write every value used to a timestamped CSV, but only if it differs from
    the newest one, so the archive is a history of real changes."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["placeholder", "value", "sheet_source"])
    for k in sorted(values):
        w.writerow([k, fmt(values[k].value), values[k].source])
    text = buf.getvalue()

    directory.mkdir(exist_ok=True)
    existing = sorted(directory.glob("freyja_params_*.csv"))
    if existing and existing[-1].read_text(encoding="utf-8") == text:
        return f"archive not updated: values identical to {existing[-1].name}"
    path = directory / f"freyja_params_{datetime.now():%Y%m%d_%H%M%S}.csv"
    path.write_text(text, encoding="utf-8")
    return f"archived values to {path.relative_to(ROOT)}"


# =============================================================================
# 6. MAIN
# =============================================================================


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Build freyja.xml from the template + Google Sheet.")
    ap.add_argument("--xlsx", metavar="FILE", help="read a downloaded copy of the sheet instead of Google")
    ap.add_argument("--strict", action="store_true", help="missing ROM values are errors, not unlimited joints")
    ap.add_argument("--dry-run", action="store_true", help="check everything but write no files")
    ap.add_argument("--no-archive", action="store_true", help="skip the CSV archive")
    ap.add_argument("--template", type=Path, default=TEMPLATE)
    ap.add_argument("--output", type=Path, default=OUTPUT)
    return ap.parse_args(argv)


def build_values(tables: dict, report: Report):
    """All parsing + data checks. Returns (values, blank_rom_names, ankle_height)."""
    bsip = read_segment_table(tables[RANGE_BSIP], BSIP_COLUMNS, report)
    pos = read_segment_table(tables[RANGE_POS], POS_COLUMNS, report)
    rom, rom_blank = read_rom_table(tables[RANGE_ROM], report)
    check_bsip(bsip, report)
    report.stop_if_errors("reading the sheet")

    values = {}
    for part in (flatten(mirror(bsip, FLIP_BSIP)), flatten(mirror(pos, FLIP_POS)), rom):
        clash = values.keys() & part.keys()
        if clash:
            report.error(f"the same name comes from two tables: {sorted(clash)}")
        values.update(part)
    ankle = check_positions(values, report)
    report.stop_if_errors("checking the sheet values")
    return values, rom_blank, ankle


def main(argv=None):
    sys.stdout.reconfigure(errors="replace")  # never crash on a odd character in the console
    args = parse_args(argv)
    report = Report()

    source = f"xlsx file {Path(args.xlsx).name}" if args.xlsx else "Google Sheet"
    print(f"[1/4] Reading {source} ...")
    tables = fetch_xlsx(args.xlsx) if args.xlsx else fetch_google()
    values, rom_blank, ankle = build_values(tables, report)
    print(f"      {sum(k.endswith('_mass') for k in values)} segments (left/right counted separately), "
          f"{len(values)} values, data checks passed")

    print(f"[2/4] Filling template {args.template.name} ...")
    if not args.template.exists():
        sys.exit(f"Template not found: {args.template}")
    text = args.template.read_text(encoding="utf-8")
    plain = {k: p.value for k, p in values.items()}
    xml = render(text, plain, report, strict=args.strict)
    report.stop_if_errors("filling the template")
    warn_unused(text, values, report)
    xml = add_banner(xml)

    print("[3/4] Checking the finished XML ...")
    check_xml(xml, report)
    model = compile_with_mujoco(xml, report) if not report.errors else None
    if report.errors:
        if not args.dry_run:
            FAILED_OUTPUT.write_text(xml, encoding="utf-8")
        report.stop_if_errors(f"checking the XML (the broken build is saved as "
                              f"{FAILED_OUTPUT.name}; {args.output.name} was NOT touched)")

    print("[4/4] Writing ...")
    notes = []
    if args.dry_run:
        notes.append("dry run: nothing written")
    else:
        old = args.output.read_text(encoding="utf-8") if args.output.exists() else None
        if old == xml:
            notes.append(f"{args.output.name} already up to date (sheet and template give the same model, nothing to write)")
        else:
            args.output.write_text(xml, encoding="utf-8")
            notes.append(f"{args.output.name} {'created' if old is None else 'updated'}")
        if not args.no_archive:
            notes.append(archive_values(values, ARCHIVE_DIR))

    print("\n" + "=" * 70)
    if model is not None:
        print(f"MuJoCo compiled OK: {model.nbody - 1} bodies, {model.njnt} joints, "
              f"total mass {sum(model.body_mass):.3f} kg")
    if ankle is not None:
        print(f"Standing pose: ankle joint {ankle * 1000:.1f} mm above the floor")
    for n in notes:
        print(n)
    if report.warnings:
        print(f"\n{len(report.warnings)} warning(s):")
        for w in report.warnings:
            print(f"  WARNING {w}")
    print("=" * 70)


if __name__ == "__main__":
    main()
